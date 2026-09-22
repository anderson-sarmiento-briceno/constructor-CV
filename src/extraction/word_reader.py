from pathlib import Path

from docx import Document


def extract_text_from_docx(docx_path):
    """Lee un archivo .docx y devuelve todo el texto plano en un solo bloque."""
    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")

    document = Document(str(path))
    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs)


def extract_keywords_from_text(text):
    """Devuelve una lista de palabras clave mínimamente útiles para matching."""
    cleaned = text.replace("\n", " ")
    tokens = []
    for chunk in cleaned.split():
        clean_chunk = chunk.strip("()[]{}.,;:-_/\"")
        if len(clean_chunk) > 2:
            tokens.append(clean_chunk)
    return tokens
