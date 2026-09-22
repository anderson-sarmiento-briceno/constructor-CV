from docx import Document

from src.engine import adapt_profile_to_offer, generate_all_cv_for_offers
from src.rendering.pdf_renderer import render_cv_to_pdf


def test_render_cv_to_pdf_creates_file(tmp_path):
    output = tmp_path / "cv_test.pdf"
    html = "<html><body><h1>CV prueba</h1><p>Texto</p></body></html>"

    path = render_cv_to_pdf(html, str(output))

    assert path == str(output)
    assert output.exists()
    assert output.stat().st_size > 0


def test_render_cv_to_pdf_accepts_structured_profile(tmp_path):
    output = tmp_path / "cv_structured.pdf"
    profile = {
        "datos_personales": {
            "nombre": "Anderson Sarmiento",
            "profesion": "Ingeniero Eléctrico",
            "correo": "anderson@email.com",
            "telefono": "+57 300 000 0000",
            "linkedin": "linkedin.com/in/anderson",
        },
        "perfil_profesional": {
            "resumen": "Ingeniero eléctrico con experiencia en automatización, análisis de datos y eficiencia energética."
        },
        "habilidades": ["Python", "Power BI", "SQL", "ETL", "Machine Learning"],
        "experiencia": [
            {"empresa": "Green Mobil", "cargo": "Científico de Datos", "descripcion": "Modelos predictivos y ETL."},
            {"empresa": "Enel Colombia", "cargo": "Ingeniero", "descripcion": "Automatización y análisis eléctrico."},
        ],
        "formacion": [
            {"titulo": "Ingeniería Eléctrica", "institucion": "Universidad Distrital", "anio": 2013}
        ],
    }

    path = render_cv_to_pdf(profile, str(output))

    assert path == str(output)
    assert output.exists()
    assert output.stat().st_size > 0


def test_generate_all_cv_for_offers_processes_batch(tmp_path):
    offers_dir = tmp_path / "ofertas"
    offers_dir.mkdir()
    profiles_dir = tmp_path / "config"
    profiles_dir.mkdir()

    doc = Document()
    doc.add_paragraph("Buscamos Data Scientist con Python, Power BI, SQL y ETL")
    doc.save(str(offers_dir / "Oferta_1.docx"))

    doc2 = Document()
    doc2.add_paragraph("Necesitamos BI con Power BI, SQL y dashboards")
    doc2.save(str(offers_dir / "Oferta_2.docx"))

    profile_path = profiles_dir / "perfil_maestro.json"
    profile_path.write_text(
        '{"datos_personales":{"nombre":"Anderson Sarmiento","profesion":"Ingeniero Eléctrico"},'
        '"perfil_profesional":{"resumen":"Ingeniero eléctrico con análisis de datos"},'
        '"habilidades":["Python","Power BI","SQL","ETL","Machine Learning","Automatización"],'
        '"experiencia":[{"empresa":"Green Mobil","cargo":"Científico de Datos","descripcion":"ETL y modelos predictivos"},{"empresa":"Enel Colombia","cargo":"Ingeniero","descripcion":"Automatización y mantenimiento"}],'
        '"formacion":[]}',
        encoding="utf-8",
    )

    results = generate_all_cv_for_offers(str(offers_dir), str(profiles_dir / "perfil_maestro.json"), str(tmp_path / "output"))

    assert len(results) == 2
    assert all("pdf" in item for item in results)
    assert any("Oferta_1" in item["pdf"] for item in results)
    assert any("Oferta_2" in item["pdf"] for item in results)


def test_adapt_profile_to_offer_prioritizes_relevant_experience():
    profile = {
        "perfil_profesional": {"resumen": "Ingeniero eléctrico"},
        "habilidades": ["Python", "Power BI", "SQL", "ETL", "Machine Learning", "ISO 50001", "Movilidad eléctrica"],
        "experiencia": [
            {"empresa": "Enel Colombia", "cargo": "Ingeniero", "descripcion": "Mantenimiento eléctrico y gestión"},
            {"empresa": "Green Mobil", "cargo": "Científico de Datos", "descripcion": "Modelos predictivos, ETL, Python, Power BI, ISO 50001 y movilidad eléctrica"},
        ],
        "formacion": [],
    }

    offer_text = "Buscamos Data Scientist con Python, Power BI, SQL, ETL, Machine Learning, ISO 50001 y movilidad eléctrica"
    adapted = adapt_profile_to_offer(profile, offer_text)

    assert adapted["keywords"][0] == "Python"
    assert adapted["experiencia"][0]["empresa"] == "Green Mobil"
    assert "ISO 50001" in adapted["keywords"]


def test_adapt_profile_to_offer_is_dynamic_for_any_offer():
    profile = {
        "perfil_profesional": {"resumen": "Ingeniero eléctrico"},
        "habilidades": ["Python", "Power BI", "SQL", "ETL", "Machine Learning", "Automatización", "Eficiencia energética"],
        "experiencia": [
            {"empresa": "Enel Colombia", "cargo": "Ingeniero", "descripcion": "Automatización y mantenimiento eléctrico"},
            {"empresa": "Green Mobil", "cargo": "Científico de Datos", "descripcion": "Modelos de consumo y BI"},
        ],
        "formacion": [],
    }

    offer_text = "Necesitamos un analista BI senior con Power BI, SQL, automatización, ETL y dashboards ejecutivos"
    adapted = adapt_profile_to_offer(profile, offer_text)

    assert "Power BI" in adapted["keywords"]
    assert "SQL" in adapted["keywords"]
    assert adapted["summary"]
    assert "BI" in adapted["summary"].upper() or "ANALISTA" in adapted["summary"].upper()
